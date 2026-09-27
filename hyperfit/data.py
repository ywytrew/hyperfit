# SPDX-License-Identifier: GPL-3.0-or-later
"""Import observations with explicit measures and units; never guess from magnitudes."""
import base64
import csv
import io
import numpy as np
import openpyxl
from .objectives import Dataset


def read_table(filename, content, sheet=None):
    raw = base64.b64decode(content, validate=True)
    if len(raw) > 10_000_000:
        raise ValueError("File exceeds the 10 MB local import limit")
    if filename.lower().endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            names = wb.sheetnames
            ws = wb[sheet or names[0]]
            if ws.max_row and ws.max_row > 50001:
                raise ValueError("Select a worksheet with at most 50000 rows")
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) > 50001: raise ValueError("Too many rows")
        finally:
            wb.close()
    elif filename.lower().endswith((".csv", ".tsv")):
        names = []
        text = raw.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text), delimiter="\t" if filename.lower().endswith(".tsv") else ","))
    else:
        raise ValueError("Use .csv, .tsv or .xlsx; convert legacy .xls explicitly first")
    if len(rows) < 2 or len(rows[0]) > 1000: raise ValueError("Missing or oversized table header")
    headers = [str(h) if h is not None else f"Column {i+1}" for i,h in enumerate(rows[0])]
    # Retain original row positions. Blank/error cells in selected columns are rejected later.
    return {"sheets": names, "headers": headers, "rows": rows[1:]}


def normalize(rows, strain_column, stress_column, volume_column, strain_measure,
              stress_measure, volume_measure, stress_unit, name="experiment"):
    if not 4 <= len(rows) <= 3000:
        raise ValueError("Select a single branch with 4–3000 points for this interactive version")
    if min(strain_column, stress_column, volume_column) < 0:
        raise ValueError("Column indices must be nonnegative")
    if len({strain_column, stress_column, volume_column}) != 3:
        raise ValueError("Choose three distinct observation columns")
    try:
        a = np.asarray([[float(row[i]) for i in (strain_column,stress_column,volume_column)] for row in rows])
    except (TypeError,ValueError,IndexError):
        raise ValueError("Selected columns contain blanks, text or non-numeric cells") from None
    if not np.all(np.isfinite(a)): raise ValueError("Non-finite observation")
    x,s,v = a.T
    with np.errstate(invalid="raise",divide="raise",over="raise"):
        if strain_measure == "engineering": x = np.log1p(x)
        elif strain_measure == "stretch": x = np.log(x)
        elif strain_measure != "log": raise ValueError("Unknown axial strain measure")
        if volume_measure == "J": v = np.log(v)
        elif volume_measure == "J_minus_1": v = np.log1p(v)
        elif volume_measure != "log_J": raise ValueError("Unknown volume measure")
        units = {"MPa": 1, "kPa": .001, "Pa": 1e-6}
        s = s*units[stress_unit]
        if stress_measure == "nominal": s = s*np.exp(x-v)
        elif stress_measure != "cauchy": raise ValueError("Only Cauchy or nominal stress is accepted")
    return Dataset(x,s,v,name)
